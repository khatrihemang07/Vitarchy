"""Spreadsheets and CSV as a Grid: rows of cells, every value as displayed text."""

from __future__ import annotations

import csv
import io

import openpyxl
import xlrd

from ..errors import OpenError
from ..source import Source, open_stream, read_bytes
from .cell_text import blank_merged_followers, datetime_text, number_text, pad_rows

Grid = list[list[str]]
DELIMITERS = ",;\t|"
UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")


def read_grid(source: Source, format_name: str) -> Grid:
    """Read the first non-empty sheet (or the CSV) of a file already identified by `sniff`."""
    readers = {"xlsx": read_xlsx, "xls": read_xls, "csv": read_csv}
    try:
        return readers[format_name](source)
    except OpenError:
        raise
    except Exception as error:
        raise OpenError("corrupt_file") from error


def _cell_text(value: object, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int | float):
        return number_text(value, number_format)
    if hasattr(value, "isoformat"):
        return datetime_text(value)  # type: ignore[arg-type]
    return str(value)


def read_xlsx(source: Source) -> Grid:
    book = openpyxl.load_workbook(open_stream(source), data_only=True)
    for sheet in book.worksheets:
        grid = [[_cell_text(c.value, c.number_format) for c in row]
                for row in sheet.iter_rows(min_row=1, min_col=1)]
        if any(cell for row in grid for cell in row):
            ranges = [(m.min_row - 1, m.max_row - 1, m.min_col - 1, m.max_col - 1)
                      for m in sheet.merged_cells.ranges]
            blank_merged_followers(grid, ranges)
            return pad_rows(grid)
    return []


def read_xls(source: Source) -> Grid:
    book = xlrd.open_workbook(file_contents=read_bytes(source), formatting_info=True)
    for sheet in book.sheets():
        grid = [[_xls_cell(book, sheet.cell(r, c)) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
        if any(cell for row in grid for cell in row):
            blank_merged_followers(
                grid, [(lo, hi - 1, clo, chi - 1) for lo, hi, clo, chi in sheet.merged_cells])
            return pad_rows(grid)
    return []


def _xls_cell(book: xlrd.Book, cell: xlrd.sheet.Cell) -> str:
    if cell.ctype == xlrd.XL_CELL_DATE:
        return datetime_text(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        number_format = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
        return number_text(cell.value, number_format)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return ""
    return str(cell.value)


def decode_text(data: bytes) -> str:
    if data.startswith(UTF16_BOMS):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def pick_delimiter(rows_by_delimiter: dict[str, list[list[str]]]) -> str:
    """The delimiter whose parse gives the most columns on the most rows of the same width."""
    def score(rows: list[list[str]]) -> int:
        widths = [len(r) for r in rows if len(r) > 1]
        if not widths:
            return 0
        modal = max(set(widths), key=widths.count)
        return widths.count(modal) * (modal - 1)

    return max(rows_by_delimiter, key=lambda d: score(rows_by_delimiter[d]))


def _parse(text: str, delimiter: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))


def read_csv(source: Source) -> Grid:
    text = decode_text(read_bytes(source))
    parsed = {d: _parse(text, d) for d in DELIMITERS}
    return pad_rows(parsed[pick_delimiter(parsed)])
