"""Spreadsheet and CSV layouts (traits j, k, l)."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time
from pathlib import Path

import msoffcrypto
import openpyxl
import xlwt

from .built import SUMMARY_KEYS, Built, expected_statements, statement
from .data import ledger, plain, totals

ISSUER = "Example Bank"


def encrypted_xlsx(out: Path) -> Built:
    """(j) Encrypted workbook: preamble, a header merged across rows and columns, a summary footer.

    Password 'XLSX2026'; the test stores 'xlsx2026', so the upper-case variant opens it.
    """
    rows = ledger(10, 10, 7_471_778, date(2026, 9, 1))
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "sheet"
    ws["B2"], ws["B3"] = f"{ISSUER}, Sample Branch", "Account Number  :  XXXXXX0789"
    ws["B4"] = "IFSC Code  :  EXMP0001234"
    ws["A5"] = "Statement From  :  01-09-2026  to  30-09-2026"
    for row in (2, 3, 4):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    ws.merge_cells("A5:F5")
    ws["A7"], ws["B7"], ws["C7"], ws["D7"], ws["F7"] = "Date", "Details", "Ref No", "Amount", "Balance"
    ws["D8"], ws["E8"] = "Debit", "Credit"
    for col in "ABCF":
        ws.merge_cells(f"{col}7:{col}8")
    ws.merge_cells("D7:E7")
    first = 9
    for i, row in enumerate(rows):
        r = first + i
        ws.cell(r, 1, datetime.combine(row.date, time()))
        ws.cell(r, 2, row.description)
        ws.cell(r, 3, "")
        ws.cell(r, 4, plain(row.debit) if row.debit else "")
        ws.cell(r, 5, plain(row.credit) if row.credit else "")
        ws.cell(r, 6, row.balance / 100).number_format = "0.00"
    foot = first + len(rows) + 1
    ws.cell(foot, 1, "Statement Summary : 01-09-2026 To 30-09-2026")
    ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=6)
    debit, credit, dcount, ccount = totals(rows)
    heads = ["Brought Forward", "Dr Count", "Cr Count", "Total Debits", "Total Credits", "Closing Balance"]
    values = [plain(7_471_778), str(dcount), str(ccount), plain(debit), plain(credit),
              plain(rows[-1].balance)]
    for c, (head, value) in enumerate(zip(heads, values, strict=True), start=1):
        ws.cell(foot + 1, c, head)
        ws.cell(foot + 2, c, value)
    path = out / "encrypted_sheet.xlsx"
    plain_bytes = io.BytesIO()
    book.save(plain_bytes)
    plain_bytes.seek(0)
    with open(path, "wb") as handle:
        msoffcrypto.OfficeFile(plain_bytes).encrypt("XLSX2026", handle)
    stmt = statement(ISSUER, "savings", "XXXXXX0789", (date(2026, 9, 1), date(2026, 9, 30)), rows,
                     opening_paise=7_471_778, closing_paise=rows[-1].balance,
                     total_debit_paise=debit, total_credit_paise=credit,
                     debit_count=dcount, credit_count=ccount)
    return Built(path.name, expected_statements([stmt]), {
        "format": "xlsx", "encrypted": True, "password": "XLSX2026", "saved_password": "xlsx2026",
        "rows": foot + 2, "cols": 6, "merged_blank": [["D7", "E7"], ["A7", "A8"]],
        "first_data_row": first,
    })


def xls_with_ghost_row(out: Path) -> Built:
    """(k) A legacy workbook: spacer columns, a merged preamble, and a remark spilling to a ghost row."""
    rows = ledger(11, 8, 8_000_000, date(2026, 7, 2))
    for i, row in enumerate(rows):
        row.description = f"UPI/{500000000000 + i}/NA/{row.description.split('/')[-1]}"
    long_tail = "2e2e47d78cb1487e1bc79e69/"
    rows[3].description = "UPI/118326283710/NA/fake@exb/Example Bank/IBL9e92608a512345712db51c8c511c8d33" \
                          + "/" + long_tail
    cap = 60
    book = xlwt.Workbook()
    sheet = book.add_sheet("History")
    sheet.write(1, 1, "DETAILED STATEMENT")
    sheet.write(3, 1, "Account Number")
    sheet.write_merge(3, 3, 3, 8, "XXXXXX0555 ( INR )  - FAKE NAME")
    sheet.write(4, 1, "Transaction Date from")
    sheet.write(4, 3, "01/07/2026")
    sheet.write(4, 4, "to")
    sheet.write(4, 5, "31/07/2026")
    sheet.write(11, 1, "Transactions List - FAKE NAME - XXXXXX0555")
    header = ["S No.", "Value Date", "Transaction Date", "Cheque Number", "Transaction Remarks",
              "Withdrawal Amount(INR)", "Deposit Amount(INR)", "Balance(INR)"]
    for c, text in enumerate(header, start=1):
        sheet.write(12, c, text)
    r = 13
    for i, row in enumerate(rows):
        spill = len(row.description) > cap
        remark = row.description[:cap] if spill else row.description
        values = [str(i + 1), row.date.strftime("%d/%m/%Y"), row.date.strftime("%d/%m/%Y"), "",
                  remark, plain(row.debit) if row.debit else "0.00",
                  plain(row.credit) if row.credit else "0.00", plain(row.balance)]
        for c, text in enumerate(values, start=1):
            sheet.write(r, c, text)
        r += 1
        if spill:
            sheet.write(r, 5, row.description[cap:])
            r += 1
    sheet.write(r + 1, 1, "Legends Used in Account Statement")
    for k, legend in enumerate(["1. INFT - Internal Fund Transfer", "2. BIL - Bill payment"]):
        sheet.write(r + 2 + k, 1, legend)
    path = out / "ghost_row.xls"
    book.save(str(path))
    stmt = statement(ISSUER, "savings", "XXXXXX0555", (date(2026, 7, 1), date(2026, 7, 31)), rows)
    return Built(path.name, expected_statements([stmt]), {
        "format": "xls", "rows": r + 4, "cols": 9, "ghost_rows": 1, "remark_cap": cap,
    })


def payment_history_csv(out: Path) -> Built:
    """(l) A Payment history file: preamble, quoted fields, 'Sept' dates, separate date and time,
    two instruments interleaved (one Statement each, in first-seen order)."""
    records = [
        ("2026-10-18T11:20:00", "Paid to FAKE SHOP", 65000, "debit", "Paid by XXXXXX0123"),
        ("2026-09-19T23:59:00", "Payment to FAKE STREAMING", 19900, "debit", "Paid by EXAMPLE R"),
        ("2026-09-30T09:05:00", "Received from FAKE PERSON", 120050, "credit", "Credited to XXXXXX0123"),
        ("2026-09-19T00:05:00", "Refund from FAKE TRAVEL", 643980, "credit", "Credited to EXAMPLE R"),
    ]
    months = {9: "Sept", 10: "Oct"}
    lines = [["Transaction Statement for 9000000000"],
             ["Duration", "18 Sept, 2026 - 18 Oct, 2026"], [],
             ["Date", "Time", "Transaction Details", "Transaction ID", "UTR", "Transaction Type",
              "Credit/debit instrument", "Amount"]]
    sections: dict[str, list[dict]] = {}
    for i, (stamp, details, paise, direction, instrument) in enumerate(records):
        moment = datetime.fromisoformat(stamp)
        txn_id = f"T{moment:%y%m%d%H%M}{i:013d}"
        utr = f"{512345678952 + i}"
        lines.append([f"{months[moment.month]} {moment.day:02d}, {moment.year}",
                      moment.strftime("%I:%M %p").lower(), details, txn_id, utr, direction.upper(),
                      instrument, plain(paise).removesuffix(".00")])
        shown = instrument.split(" ", 2)[-1]
        sections.setdefault(shown, []).append({
            "date": moment.date().isoformat(), "time": moment.strftime("%H:%M:%S"),
            "description": details, "amount_paise": paise, "direction": direction,
            "balance_paise": None, "foreign_amount": None, "foreign_currency": None,
            "reference": utr, "source_id": txn_id})
    statements = []
    for shown, found in sections.items():
        days = [t["date"] for t in found]
        statements.append({
            "account": {"issuer": ISSUER if shown == "EXAMPLE R" else None, "kind": None, "mask": shown},
            "period": {"start": min(days), "end": max(days)},
            "summary": dict.fromkeys(SUMMARY_KEYS), "transactions": found})
    lines += [[], ["This is an automatically generated statement."], ["Disclaimer : fake text", "x"]]
    path = out / "payment_history.csv"
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, quoting=csv.QUOTE_ALL, lineterminator="\n")
        for i, line in enumerate(lines):
            if i == 3:  # the header row is not quoted
                handle.write(",".join(line) + "\n")
            else:
                writer.writerow(line)
    return Built(path.name, {"outcome": "read", "kind": "payment_history",
                             "statements": statements},
                 {"format": "csv", "rows": len(lines), "cols": 8, "bom": True})
